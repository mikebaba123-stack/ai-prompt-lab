"""สร้างไฟล์ Excel พร้อมกราฟ insight จากข้อมูลที่ทำความสะอาดแล้ว (เปิดใน Google Sheets ได้)"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

SRC = "Social_Media_Usage_Survey_clean.csv"
OUT = "insight_social_media.xlsx"
TIME = ["Facebook_Time","Instagram_Time","TikTok_Time","YouTube_Time"]

df = pd.read_csv(SRC)
wb = Workbook(); wb.remove(wb.active)

HDR = PatternFill("solid", fgColor="06192F")
HFONT = Font(color="FFFFFF", bold=True)

def sheet(name, frame, number_format=None):
    ws = wb.create_sheet(name)
    for r in dataframe_to_rows(frame, index=False, header=True):
        ws.append(r)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HDR, HFONT, Alignment(horizontal="center")
    for col in ws.columns:
        width = max(len(str(c.value)) for c in col if c.value is not None) + 3
        ws.column_dimensions[col[0].column_letter].width = min(width, 26)
    if number_format:
        for row in ws.iter_rows(min_row=2, min_col=2):
            for c in row: c.number_format = number_format
    return ws

# ── ชีต 1 · ข้อมูลสะอาดทั้งหมด
sheet("ข้อมูลสะอาด", df)

# ── ชีต 2 · เวลาเฉลี่ยรายแพลตฟอร์ม แยกตามช่วงอายุ  → กราฟแท่ง
by_age = df.groupby("Age_Group", observed=True)[TIME].mean().round(2).reset_index()
ws = sheet("เวลาเฉลี่ยตามช่วงอายุ", by_age, "0.00")
ch = BarChart(); ch.type="col"; ch.grouping="clustered"
ch.title = "เวลาใช้งานเฉลี่ยต่อวัน (ชม.) แยกตามช่วงอายุและแพลตฟอร์ม"
ch.y_axis.title = "ชั่วโมง/วัน"; ch.x_axis.title = "ช่วงอายุ"
ch.add_data(Reference(ws, min_col=2, max_col=5, min_row=1, max_row=ws.max_row), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
ch.height, ch.width = 9, 20
ws.add_chart(ch, "H2")

# ── ชีต 3 · สัดส่วนเวลาทั้งหมดรายแพลตฟอร์ม → กราฟวงกลม
share = pd.DataFrame({"แพลตฟอร์ม":[c.replace("_Time","") for c in TIME],
                      "ชั่วโมงรวม":[round(df[c].sum(),1) for c in TIME]})
ws = sheet("สัดส่วนรายแพลตฟอร์ม", share, "0.0")
pie = PieChart(); pie.title = "สัดส่วนเวลาใช้งานรวมของกลุ่มตัวอย่าง 150 คน"
pie.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
pie.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
pie.height, pie.width = 9, 14
ws.add_chart(pie, "E2")

# ── ชีต 4 · คุณภาพการนอน เทียบกับเวลาใช้งานรวม → กราฟแท่ง
sleep = (df[df["Sleep_Quality"]!="Unknown"]
         .groupby("Sleep_Quality")["Total_Time_Hours"]
         .agg(["mean","count"]).round(2)
         .rename(columns={"mean":"เวลาใช้งานเฉลี่ย (ชม.)","count":"จำนวนคน"})
         .reindex(["Good","Fair","Poor"]).reset_index())
ws = sheet("นอนหลับ vs เวลาใช้งาน", sleep, "0.00")
ch2 = BarChart(); ch2.type="col"
ch2.title = "เวลาใช้โซเชียลเฉลี่ยต่อวัน แยกตามคุณภาพการนอน"
ch2.y_axis.title = "ชั่วโมง/วัน"; ch2.x_axis.title = "คุณภาพการนอน"
ch2.add_data(Reference(ws, min_col=2, max_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
ch2.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
ch2.height, ch2.width = 9, 16
ws.add_chart(ch2, "F2")

# ── ชีต 5 · อารมณ์ เทียบกับเวลาใช้งาน
emo = (df[df["Emotion"]!="Unknown"]
       .groupby("Emotion")["Total_Time_Hours"].agg(["mean","count"]).round(2)
       .rename(columns={"mean":"เวลาใช้งานเฉลี่ย (ชม.)","count":"จำนวนคน"}).reset_index())
ws = sheet("อารมณ์ vs เวลาใช้งาน", emo, "0.00")
ch3 = BarChart(); ch3.type="bar"
ch3.title = "เวลาใช้โซเชียลเฉลี่ยต่อวัน แยกตามอารมณ์ที่รายงาน"
ch3.add_data(Reference(ws, min_col=2, max_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
ch3.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
ch3.height, ch3.width = 9, 16
ws.add_chart(ch3, "F2")

wb.save(OUT)
print("บันทึก:", OUT)
print("\n--- ตัวเลขที่ได้ ---")
print(by_age.to_string(index=False)); print()
print(share.to_string(index=False)); print()
print(sleep.to_string(index=False)); print()
print(emo.to_string(index=False))
